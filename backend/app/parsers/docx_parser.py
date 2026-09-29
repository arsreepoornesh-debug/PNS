import docx
import logging
from app.parsers.base import ParsedDocument, DocumentPage
from app.parsers.pdf_parser import clean_text

logger = logging.getLogger(__name__)

def parse_docx(file_path: str) -> ParsedDocument:
    """
    Extracts structured paragraphs and tables from DOCX documents.
    """
    try:
        doc = docx.Document(file_path)
        paragraphs_text = []
        
        for p in doc.paragraphs:
            text = clean_text(p.text)
            if text:
                paragraphs_text.append(text)
                
        # Also extract table text
        for table in doc.tables:
            for row in table.rows:
                row_cells = [clean_text(cell.text) for cell in row.cells if clean_text(cell.text)]
                if row_cells:
                    paragraphs_text.append(" | ".join(row_cells))
                    
        full_text = "\n\n".join(paragraphs_text).strip()
        is_empty = len(full_text) < 20
        
        pages = [DocumentPage(page_number=1, text=full_text)]
        
        return ParsedDocument(
            raw_text=full_text,
            pages=pages,
            page_count=1,
            file_type="docx",
            is_scanned_or_empty=is_empty,
            warning="Document is empty." if is_empty else None
        )
    except Exception as e:
        logger.error(f"Error parsing DOCX {file_path}: {e}", exc_info=True)
        return ParsedDocument(
            raw_text="",
            pages=[],
            page_count=1,
            file_type="docx",
            is_scanned_or_empty=True,
            warning=f"Malformed or unreadable DOCX: {str(e)}"
        )
